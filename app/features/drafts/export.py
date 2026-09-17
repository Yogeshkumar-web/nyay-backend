"""
Synchronous PDF and DOCX generation from draft content.
Handles raw text/markdown and legacy HTML, then applies code-controlled formatting.
Used by the direct download endpoint — no Celery, no R2.
"""
import io
import re
from typing import Any
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from xhtml2pdf import pisa


# ============================================================
# MARKDOWN → HTML CONVERTER
# ============================================================


def _inline_md(text: str) -> str:
    """Convert inline markdown (bold, italic) to HTML."""
    # Bold + italic
    text = re.sub(r"\*\*\*(.+?)\*\*\*", r"<strong><em>\1</em></strong>", text)
    # Bold
    text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
    # Italic (single asterisk, not greedy across lines)
    text = re.sub(r"\*([^*\n]+?)\*", r"<em>\1</em>", text)
    return text


def _md_to_html(text: str) -> str:
    """
    Lightweight markdown → HTML for AI-generated legal document patterns.
    Handles: headings, bold/italic, horizontal rules, paragraphs, lists.
    Falls through immediately if content already contains HTML tags.
    """
    if re.search(r"<[a-zA-Z][^>]*>", text):
        return text  # already HTML from legacy saved content

    html_parts: list[str] = []
    # Normalise line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Split into blocks on blank lines
    blocks = re.split(r"\n{2,}", text.strip())

    for block in blocks:
        block = block.strip()
        if not block:
            continue

        # Horizontal rule
        if re.match(r"^[-*_]{3,}$", block):
            html_parts.append("<hr/>")
            continue

        # Heading
        m = re.match(r"^(#{1,6})\s+(.*)", block)
        if m:
            level = len(m.group(1))
            content = _inline_md(m.group(2).strip())
            html_parts.append(f"<h{level}>{content}</h{level}>")
            continue

        # Unordered list (lines starting with - or *)
        lines = block.split("\n")
        if all(re.match(r"^[-*]\s+", line.strip()) for line in lines if line.strip()):
            items = "".join(
                f'<li>{_inline_md(re.sub(r"^[-*]\s+", "", line.strip()))}</li>'
                for line in lines
                if line.strip()
            )
            html_parts.append(f"<ul>{items}</ul>")
            continue

        # Ordered list
        if all(
            re.match(r"^\d+[.)]\s+", line.strip()) for line in lines if line.strip()
        ):
            items = "".join(
                f'<li>{_inline_md(re.sub(r"^\d+[.)]\s+", "", line.strip()))}</li>'
                for line in lines
                if line.strip()
            )
            html_parts.append(f"<ol>{items}</ol>")
            continue

        # Plain paragraph — join lines within the block with a space
        combined = " ".join(line.strip() for line in lines if line.strip())
        html_parts.append(f"<p>{_inline_md(combined)}</p>")

    return "\n".join(html_parts)


# ============================================================
# PDF/DOCX CSS
# ============================================================

PAGE_CONFIG = {
    "size": {"width": 12240, "height": 20160},
    "margin": {"top": 2160, "bottom": 1440, "left": 2520, "right": 720},
}


_LEGAL_CSS = """
@page {
    size: legal;
    margin-left: 1.75in;
    margin-right: 0.5in;
    margin-top: 1.5in;
    margin-bottom: 1in;
}
body {
    font-family: "Times New Roman", Times, serif;
    font-size: 13pt;
    line-height: 1.8;
    color: #000;
}
p {
    text-align: justify;
    margin-bottom: 8pt;
    orphans: 3;
    widows: 3;
}
h1, h2, h3 {
    font-weight: bold;
    text-align: center;
    margin-top: 14pt;
    margin-bottom: 6pt;
}
h1 { font-size: 14pt; text-transform: uppercase; }
h2 { font-size: 13pt; text-transform: uppercase; }
h3 { font-size: 13pt; }
ul, ol { margin-left: 1.5cm; margin-bottom: 8pt; }
li { margin-bottom: 3pt; }
strong { font-weight: bold; }
em { font-style: italic; }
hr { border: none; border-top: 1px solid #000; margin: 10pt 0; }
table {
    width: 100%;
    border-collapse: collapse;
    table-layout: fixed;
    margin: 8pt 0 12pt;
}
th, td {
    border: 1px solid #555;
    padding: 5pt 6pt;
    vertical-align: top;
    line-height: 1.35;
    word-wrap: break-word;
}
th {
    font-weight: bold;
    background-color: #f5f5f5;
}
"""


# ============================================================
# PUBLIC API
# ============================================================


def generate_pdf(content: str, title: str) -> bytes:
    """Convert reviewed text/markdown/legacy HTML to a code-formatted PDF."""
    html_body = _md_to_html(content)
    styled = (
        f"<html><head>"
        f'<meta charset="utf-8"/>'
        f"<style>{_LEGAL_CSS}</style>"
        f"</head><body>{html_body}</body></html>"
    )
    buf = io.BytesIO()
    pisa.CreatePDF(styled, dest=buf, encoding="utf-8")
    return buf.getvalue()


def generate_docx(content: str, title: str) -> bytes:
    """Convert reviewed text/markdown/legacy HTML to a code-formatted DOCX."""
    return generate_reviewed_document_docx(content, title)


def generate_reviewed_document_docx(
    content: str,
    title: str,
    *,
    structured_extraction: dict[str, Any] | None = None,
    canonical_document: dict[str, Any] | None = None,
) -> bytes:
    """Generate a deterministic DOCX from lawyer-reviewed text and extraction hints."""
    html_body = _md_to_html(content)

    doc = Document()

    # Legal 8.5x14 page setup. Values are kept in twips to match the frontend
    # review surface and avoid cm/in rounding drift.
    section = doc.sections[0]
    section.page_width = Twips(PAGE_CONFIG["size"]["width"])
    section.page_height = Twips(PAGE_CONFIG["size"]["height"])
    section.top_margin = Twips(PAGE_CONFIG["margin"]["top"])
    section.bottom_margin = Twips(PAGE_CONFIG["margin"]["bottom"])
    section.left_margin = Twips(PAGE_CONFIG["margin"]["left"])
    section.right_margin = Twips(PAGE_CONFIG["margin"]["right"])

    # Default paragraph style — Times New Roman 13pt, 1.8 line spacing
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    para_fmt = style.paragraph_format
    para_fmt.space_after = Pt(6)
    para_fmt.line_spacing = Pt(23)

    _configure_named_styles(doc)

    if canonical_document:
        from app.features.documents.canonical_document import (
            canonical_document_from_dict,
        )

        canonical = canonical_document_from_dict(canonical_document)
        if canonical is not None:
            _add_document_title(doc, canonical.title or title)
            _add_canonical_blocks(doc, canonical.blocks)
            buf = io.BytesIO()
            doc.save(buf)
            return buf.getvalue()

    _add_document_title(doc, title)

    soup = BeautifulSoup(html_body, "html.parser")
    body = soup.body or soup

    def _apply_inline(run, tag):
        """Walk ancestors to apply bold/italic/underline."""
        for parent in tag.parents:
            if parent.name in ("strong", "b"):
                run.bold = True
            if parent.name in ("em", "i"):
                run.italic = True
            if parent.name == "u":
                run.underline = True

    def _add_para(
        el,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        bold_all=False,
        font_size: int | None = None,
    ):
        p = doc.add_paragraph(style="Normal")
        p.alignment = alignment
        for child in el.children:
            if hasattr(child, "name") and child.name:
                run = p.add_run(child.get_text())
                _apply_inline(run, child)
            else:
                run = p.add_run(str(child))
            if bold_all:
                run.bold = True
            if font_size:
                run.font.size = Pt(font_size)
        return p

    for el in body.children:
        if not hasattr(el, "name") or not el.name:
            continue
        tag = el.name
        text = el.get_text().strip()
        if not text:
            continue

        if tag == "h1":
            _add_para(
                el, alignment=WD_ALIGN_PARAGRAPH.CENTER, bold_all=True, font_size=14
            )
        elif tag == "h2":
            _add_para(
                el, alignment=WD_ALIGN_PARAGRAPH.CENTER, bold_all=True, font_size=13
            )
        elif tag == "h3":
            _add_para(el, bold_all=True)
        elif tag == "hr":
            # Thin horizontal rule via bottom border on an empty paragraph
            p = doc.add_paragraph()
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement

            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "000000")
            pBdr.append(bottom)
            pPr.append(pBdr)
        elif tag in ("ul", "ol"):
            for i, li in enumerate(el.find_all("li", recursive=False), start=1):
                bullet = "•" if tag == "ul" else f"{i}."
                p = doc.add_paragraph(style="Normal")
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.left_indent = Twips(720)
                p.paragraph_format.first_line_indent = Twips(-360)
                p.add_run(f"{bullet}  {li.get_text().strip()}")
        else:
            _add_reviewed_paragraphs(doc, el, _add_para)

    extraction_result = _structured_result(structured_extraction)
    if extraction_result:
        _add_structured_appendix(doc, extraction_result)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_canonical_blocks(doc: Document, blocks: list[Any]) -> None:
    for block in blocks:
        if block.semantic_role == "title" and block.type in {"heading", "generic"}:
            continue
        if block.type in {"heading", "generic"} and block.heading:
            paragraph = doc.add_paragraph(style="Heading 2")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.add_run(block.heading)
        elif block.type in {"paragraph", "signature"} and block.text:
            paragraph = doc.add_paragraph(style="Normal")
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
                if block.type == "signature"
                else WD_ALIGN_PARAGRAPH.JUSTIFY
            )
            paragraph.add_run(block.text)
        elif block.type == "key_value" and block.fields:
            table = doc.add_table(rows=0, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for field in block.fields:
                cells = table.add_row().cells
                cells[0].text = field.label
                cells[1].text = field.value
                for run in cells[0].paragraphs[0].runs:
                    run.bold = True
        elif block.type == "table" and (block.headers or block.rows):
            column_count = max(
                len(block.headers),
                max((len(row) for row in block.rows), default=0),
            )
            if column_count == 0:
                continue
            table = doc.add_table(rows=0, cols=column_count)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            if block.headers:
                cells = table.add_row().cells
                for index, value in enumerate(block.headers):
                    cells[index].text = value
                    for run in cells[index].paragraphs[0].runs:
                        run.bold = True
            for row in block.rows:
                cells = table.add_row().cells
                for index, value in enumerate(row[:column_count]):
                    cells[index].text = value
        elif block.type == "list" and block.items:
            for index, item in enumerate(block.items, start=1):
                paragraph = doc.add_paragraph(style="Normal")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                paragraph.paragraph_format.left_indent = Twips(720)
                paragraph.paragraph_format.first_line_indent = Twips(-360)
                marker = f"{index}." if block.ordered else "•"
                paragraph.add_run(f"{marker}  {item}")
        elif block.type == "page_break":
            doc.add_page_break()


def _configure_named_styles(doc: Document) -> None:
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
    for style_name, size in (("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 13)):
        style = doc.styles[style_name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)


def _add_document_title(doc: Document, title: str) -> None:
    safe_title = (title or "Reviewed Document").strip()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(safe_title.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)


def _add_reviewed_paragraphs(doc: Document, el, add_para) -> None:
    text = el.get_text("\n").strip()
    page_blocks = _split_page_marked_text(text)
    if not page_blocks:
        add_para(el)
        return
    for block in page_blocks:
        if block["page_number"] is not None:
            note = doc.add_paragraph(style="Normal")
            note.alignment = WD_ALIGN_PARAGRAPH.LEFT
            note.paragraph_format.space_before = Pt(4)
            note.paragraph_format.space_after = Pt(2)
            run = note.add_run(f"Source page {block['page_number']}")
            run.italic = True
            run.font.size = Pt(10)
        for paragraph_text in re.split(r"\n{2,}", block["text"].strip()):
            if not paragraph_text.strip():
                continue
            paragraph = doc.add_paragraph(style="Normal")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.add_run(" ".join(line.strip() for line in paragraph_text.splitlines()))


def _split_page_marked_text(text: str) -> list[dict[str, Any]]:
    marker_re = re.compile(
        r"\[\[PAGE\s+(\d+)\s+START\]\](.*?)\[\[PAGE\s+\1\s+END\]\]",
        re.IGNORECASE | re.DOTALL,
    )
    matches = list(marker_re.finditer(text))
    if not matches:
        return []
    blocks: list[dict[str, Any]] = []
    for match in matches:
        blocks.append(
            {
                "page_number": int(match.group(1)),
                "text": match.group(2).strip(),
            }
        )
    return blocks


def _structured_result(structured_extraction: dict[str, Any] | None) -> dict[str, Any]:
    if not structured_extraction:
        return {}
    result = structured_extraction.get("result", structured_extraction)
    return result if isinstance(result, dict) else {}


def _add_structured_appendix(doc: Document, result: dict[str, Any]) -> None:
    dates = result.get("dates") if isinstance(result.get("dates"), list) else []
    names = result.get("names") if isinstance(result.get("names"), list) else []
    warnings = result.get("warnings") if isinstance(result.get("warnings"), list) else []
    unclear = (
        result.get("unclear_words") if isinstance(result.get("unclear_words"), list) else []
    )
    if not dates and not names and not warnings and not unclear:
        return

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("REVIEW TRACE")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)

    if names:
        _add_table(
            doc,
            "Names flagged during extraction",
            ["Name", "Role", "Source pages", "Warning"],
            [
                [
                    str(item.get("raw_text") or ""),
                    str(item.get("role") or ""),
                    _pages_text(item.get("source_pages")),
                    str(item.get("warning") or ""),
                ]
                for item in names
                if isinstance(item, dict)
            ],
        )
    if dates:
        _add_table(
            doc,
            "Dates flagged during extraction",
            ["Raw date", "Normalized", "Context", "Source pages"],
            [
                [
                    str(item.get("raw_text") or ""),
                    str(item.get("normalized_date") or ""),
                    str(item.get("context") or ""),
                    _pages_text(item.get("source_pages")),
                ]
                for item in dates
                if isinstance(item, dict)
            ],
        )
    if unclear or warnings:
        _add_table(
            doc,
            "Review flags",
            ["Type", "Value"],
            [["Unclear word", str(word)] for word in unclear]
            + [["Warning", str(warning)] for warning in warnings],
        )


def _add_table(
    doc: Document,
    title: str,
    headers: list[str],
    rows: list[list[str]],
) -> None:
    if not rows:
        return
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(12)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)


def _pages_text(value: Any) -> str:
    if not isinstance(value, list):
        return ""
    return ", ".join(str(page) for page in value)
