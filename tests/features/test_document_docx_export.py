from __future__ import annotations

import io
import uuid

import pytest
from docx import Document

import app.db.registry  # noqa: F401
from app.core.exceptions import ValidationError
from app.features.drafts.export import PAGE_CONFIG, generate_reviewed_document_docx
from app.features.documents.router import _require_reviewed_content_for_docx
from app.features.documents.repository import DocumentRepository


def _open_docx(content: bytes):
    return Document(io.BytesIO(content))


def test_reviewed_docx_uses_legal_page_settings():
    doc = _open_docx(generate_reviewed_document_docx("Reviewed paragraph", "Statement"))
    section = doc.sections[0]

    assert section.page_width.twips == PAGE_CONFIG["size"]["width"]
    assert section.page_height.twips == PAGE_CONFIG["size"]["height"]
    assert section.top_margin.twips == PAGE_CONFIG["margin"]["top"]
    assert section.bottom_margin.twips == PAGE_CONFIG["margin"]["bottom"]
    assert section.left_margin.twips == PAGE_CONFIG["margin"]["left"]
    assert section.right_margin.twips == PAGE_CONFIG["margin"]["right"]
    assert doc.styles["Normal"].font.name == "Times New Roman"


def test_docx_generation_is_blocked_until_review_content_exists():
    with pytest.raises(ValidationError, match="Save reviewed text"):
        _require_reviewed_content_for_docx(None)
    with pytest.raises(ValidationError, match="Save reviewed text"):
        _require_reviewed_content_for_docx("   ")

    assert _require_reviewed_content_for_docx("Reviewed text") == "Reviewed text"


def test_reviewed_docx_converts_page_markers_to_source_notes():
    content = (
        "[[PAGE 1 START]]\n"
        "First page text\n"
        "[[PAGE 1 END]]\n\n"
        "[[PAGE 2 START]]\n"
        "Second page text\n"
        "[[PAGE 2 END]]"
    )

    doc = _open_docx(generate_reviewed_document_docx(content, "Victim Statement"))
    texts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text]

    assert "Source page 1" in texts
    assert "First page text" in texts
    assert "Source page 2" in texts
    assert "Second page text" in texts
    assert "[[PAGE 1 START]]" not in "\n".join(texts)


def test_reviewed_docx_appends_structured_review_trace():
    structured = {
        "result": {
            "names": [
                {"raw_text": "A", "role": "victim", "source_pages": [1]},
            ],
            "dates": [
                {
                    "raw_text": "10/01/2026",
                    "normalized_date": "2026-01-10",
                    "context": "incident date",
                    "source_pages": [1],
                }
            ],
            "unclear_words": ["unclear"],
            "warnings": ["verify spelling"],
        }
    }

    doc = _open_docx(
        generate_reviewed_document_docx(
            "Reviewed paragraph",
            "Statement",
            structured_extraction=structured,
        )
    )
    paragraph_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )

    assert "REVIEW TRACE" in paragraph_text
    assert "Names flagged during extraction" in paragraph_text
    assert "Dates flagged during extraction" in paragraph_text
    assert "A" in table_text
    assert "10/01/2026" in table_text
    assert "verify spelling" in table_text


@pytest.mark.asyncio
async def test_record_docx_export_stores_metadata_only():
    session = _FakeSession()
    repo = DocumentRepository(session)  # type: ignore[arg-type]
    doc = _FakeDocument()
    exporter_id = uuid.uuid4()

    export = await repo.record_docx_export(
        doc,  # type: ignore[arg-type]
        exported_by=exporter_id,
        filename="statement.docx",
        size_bytes=1234,
        checksum_sha256="a" * 64,
        metadata={"source": "test"},
    )

    assert export.document_id == doc.id
    assert export.case_id == doc.case_id
    assert export.exported_by == exporter_id
    assert export.filename == "statement.docx"
    assert export.size_bytes == 1234
    assert export.checksum_sha256 == "a" * 64
    assert export.export_metadata == {"source": "test"}
    assert not hasattr(export, "r2_key")
    assert not hasattr(export, "r2_bucket")
    assert session.added == [export]
    assert session.flush_count == 1


class _FakeDocument:
    def __init__(self):
        self.id = uuid.uuid4()
        self.case_id = uuid.uuid4()


class _FakeSession:
    def __init__(self):
        self.added = []
        self.flush_count = 0

    def add(self, value):
        self.added.append(value)

    async def flush(self):
        self.flush_count += 1
