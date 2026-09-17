import io

import fitz
import pytest
from docx import Document as DocxDocument

from app.features.documents.digital_extractor import (
    DOCX_MIME,
    PDF_MIME,
    DocumentInputError,
    build_pdf_chunks,
    extract_docx,
    embedded_text_quality,
    is_embedded_text_usable,
    validate_file_signature,
)
from app.features.documents.document_classifier import classify_document
from app.features.documents.models import ProcessingRoute
from app.features.documents.document_processing_service import (
    OcrResult,
    process_document_bytes,
)


def _pdf_bytes(*page_texts: str) -> bytes:
    pdf = fitz.open()
    try:
        for text in page_texts:
            page = pdf.new_page()
            if text == "[[SCANNED_PAGE]]":
                page.draw_rect(fitz.Rect(72, 72, 320, 520), color=(0, 0, 0))
            elif text:
                page.insert_text((72, 72), text)
        return pdf.tobytes()
    finally:
        pdf.close()


def _docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_heading("Bail Application", level=1)
    document.add_paragraph("This is the first paragraph.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Case"
    table.rows[0].cells[1].text = "123/2026"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_classifies_digital_pdf() -> None:
    result = classify_document(_pdf_bytes("A" * 80), PDF_MIME)

    assert result.route == ProcessingRoute.digital_extract
    assert result.pdf is not None
    assert result.pdf.digital_pages == (1,)
    assert result.pdf.scanned_pages == ()
    assert result.pdf.low_quality_pages == ()


def test_rejects_mojibake_embedded_text_as_digital_quality() -> None:
    corrupted = (
        "P.S. (थाना): ठȡ\u0080Ǖ\u009aɮ\u009dȡ\u009aȡ "
        "FIR Ĥ.सǗ.ǐ\u009a. सं. 0059 "
        "भȡ\u009aतȢ\u0099 Û\u0099ȡ\u0099  ंǑ¡तȡ"
    )

    quality = embedded_text_quality(corrupted)

    assert quality["control_chars"] >= 2
    assert quality["usable"] is False
    assert is_embedded_text_usable(corrupted) is False
    assert is_embedded_text_usable(
        "प्रथम सूचना रिपोर्ट थाना ठाकुरद्वारा FIR No. 0059"
    ) is True


def test_classifies_hybrid_pdf_page_by_page() -> None:
    result = classify_document(
        _pdf_bytes("A" * 80, "[[SCANNED_PAGE]]"),
        PDF_MIME,
    )

    assert result.route == ProcessingRoute.hybrid_extract
    assert result.pdf is not None
    assert result.pdf.digital_pages == (1,)
    assert result.pdf.scanned_pages == (2,)


def test_classifies_image_as_scanned() -> None:
    result = classify_document(b"\x89PNG\r\n\x1a\nrest", "image/png")

    assert result.route == ProcessingRoute.scanned_ocr


def test_rejects_mismatched_signature() -> None:
    with pytest.raises(DocumentInputError, match="does not match"):
        validate_file_signature(b"not a pdf", PDF_MIME)


def test_extracts_docx_paragraphs_and_tables() -> None:
    text, artifact = extract_docx(_docx_bytes())

    assert "Bail Application" in text
    assert "Case | 123/2026" in text
    assert [block["type"] for block in artifact["blocks"]] == [
        "paragraph",
        "paragraph",
        "table",
    ]


def test_build_pdf_chunks_preserves_original_page_numbers() -> None:
    chunks = build_pdf_chunks(_pdf_bytes("", "", ""), (1, 3), max_pages=1)

    assert [page_numbers for page_numbers, _ in chunks] == [(1,), (3,)]


def test_rejects_password_protected_pdf() -> None:
    pdf = fitz.open()
    try:
        pdf.new_page()
        file_bytes = pdf.tobytes(
            encryption=fitz.PDF_ENCRYPT_AES_256,
            owner_pw="owner",
            user_pw="user",
        )
    finally:
        pdf.close()

    with pytest.raises(DocumentInputError, match="Password-protected"):
        classify_document(file_bytes, PDF_MIME)


def test_docx_signature_validation() -> None:
    validate_file_signature(_docx_bytes(), DOCX_MIME)


class _FakeOcrService:
    def __init__(self):
        self.calls: list[tuple[str, tuple[int, ...] | None]] = []

    async def process(
        self,
        _file_bytes: bytes,
        mime_type: str,
        *,
        original_page_numbers: tuple[int, ...] | None = None,
    ) -> OcrResult:
        self.calls.append((mime_type, original_page_numbers))
        page_numbers = original_page_numbers or (1,)
        pages = [
            {
                "page_number": page_number,
                "source": "fake_ocr",
                "text": f"OCR page {page_number}",
            }
            for page_number in page_numbers
        ]
        return OcrResult(
            text="\n\n".join(page["text"] for page in pages),
            language="hi",
            page_count=len(pages),
            artifact={"pages": pages},
        )


@pytest.mark.asyncio
async def test_hybrid_processing_only_sends_scanned_pages_to_ocr() -> None:
    service = _FakeOcrService()

    result = await process_document_bytes(
        _pdf_bytes("A" * 80, "[[SCANNED_PAGE]]", "B" * 80),
        PDF_MIME,
        ocr_service=service,
    )

    assert result.route == ProcessingRoute.hybrid_extract
    assert service.calls == [(PDF_MIME, (2,))]
    assert result.source_artifact["pages"][1]["text"] == "OCR page 2"
    assert result.ocr_language == "hi"


@pytest.mark.asyncio
async def test_digital_docx_does_not_call_ocr() -> None:
    service = _FakeOcrService()

    result = await process_document_bytes(
        _docx_bytes(),
        DOCX_MIME,
        ocr_service=service,
    )

    assert result.route == ProcessingRoute.digital_extract
    assert service.calls == []
    assert "Bail Application" in result.source_text
